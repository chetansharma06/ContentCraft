import csv
import json
import math
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from sqlmodel import Session

from app.db.models import Product
from app.db.session import create_db, engine


@dataclass(frozen=True)
class RejectedRow:
    row_number: int
    data: dict[str, Any]
    reasons: list[str]


@dataclass
class IngestionResult:
    valid_rows: list[Product]
    rejected_rows: list[RejectedRow]
    inserted_count: int


def _validate_rows(
    headers: list[Any], records: Iterable[list[Any]]
) -> tuple[list[Product], list[RejectedRow]]:
    normalized_headers = [str(header).strip().lower() if header is not None else "" for header in headers]
    missing_headers = {"name", "price"} - set(normalized_headers)
    if missing_headers:
        missing = ", ".join(sorted(missing_headers))
        raise ValueError(f"Dataset is missing required columns: {missing}")

    valid_rows: list[Product] = []
    rejected_rows: list[RejectedRow] = []
    for row_number, values in enumerate(records, start=2):
        if not any(value not in (None, "") for value in values):
            continue

        row = {
            header: values[index] if index < len(values) else None
            for index, header in enumerate(normalized_headers)
            if header
        }
        reasons: list[str] = []

        raw_name = row.get("name")
        name = str(raw_name).strip() if raw_name is not None else ""
        if not name:
            reasons.append("Product name is required")

        raw_price = row.get("price")
        try:
            if isinstance(raw_price, bool) or raw_price is None:
                raise InvalidOperation
            price_decimal = Decimal(str(raw_price).strip())
            price = float(price_decimal)
            if not price_decimal.is_finite() or price_decimal < 0 or not math.isfinite(price):
                raise InvalidOperation
        except (InvalidOperation, TypeError, ValueError, OverflowError):
            price = 0.0
            reasons.append("Price must be a finite non-negative number")

        raw_specifications = row.get("specifications")
        if raw_specifications in (None, ""):
            specifications: dict[str, Any] = {}
        else:
            try:
                specifications = (
                    raw_specifications
                    if isinstance(raw_specifications, dict)
                    else json.loads(str(raw_specifications))
                )
                if not isinstance(specifications, dict):
                    raise ValueError
            except (json.JSONDecodeError, TypeError, ValueError):
                specifications = {}
                reasons.append("Specifications must contain a valid JSON object")

        raw_keywords = row.get("keywords")
        keywords = (
            [keyword.strip() for keyword in str(raw_keywords).split(";") if keyword.strip()]
            if raw_keywords not in (None, "")
            else []
        )

        if reasons:
            rejected_rows.append(
                RejectedRow(row_number=row_number, data=row, reasons=reasons)
            )
            continue

        raw_description = row.get("description")
        description = str(raw_description).strip() if raw_description is not None else ""
        valid_rows.append(
            Product(
                name=name,
                price=price,
                description=description,
                specifications=specifications,
                keywords=keywords,
            )
        )

    return valid_rows, rejected_rows


def load_products(path: str | Path) -> tuple[list[Product], list[RejectedRow]]:
    """Parse a CSV or XLSX file into valid products and row-level rejections."""
    source = Path(path)
    suffix = source.suffix.lower()

    if suffix == ".csv":
        with source.open("r", encoding="utf-8-sig", newline="") as file:
            reader = csv.reader(file)
            headers = next(reader, None)
            if headers is None:
                raise ValueError("Dataset is empty")
            return _validate_rows(headers, reader)

    if suffix == ".xlsx":
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            worksheet = workbook.active
            records = worksheet.iter_rows(values_only=True)
            headers = next(records, None)
            if headers is None:
                raise ValueError("Dataset is empty")
            return _validate_rows(list(headers), records)
        finally:
            workbook.close()

    raise ValueError(f"Unsupported dataset format: {suffix or '(no extension)'}")


def _store_products(products: list[Product], session: Session) -> None:
    try:
        session.add_all(products)
        session.commit()
        for product in products:
            session.refresh(product)
    except Exception:
        session.rollback()
        raise


def ingest_products(
    path: str | Path,
    session: Session | None = None,
) -> IngestionResult:
    """Load a dataset and persist only its valid product rows."""
    valid_rows, rejected_rows = load_products(path)

    if session is None:
        create_db()
        with Session(engine) as database_session:
            _store_products(valid_rows, database_session)
    else:
        _store_products(valid_rows, session)

    return IngestionResult(
        valid_rows=valid_rows,
        rejected_rows=rejected_rows,
        inserted_count=len(valid_rows),
    )